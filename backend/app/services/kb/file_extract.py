"""
统一文件解析：PDF / Word / Excel / CSV / TXT → 纯文本 + Excel 结构化行（如有）。
"""
from __future__ import annotations

import csv
import io
import logging
import os
import re
from dataclasses import dataclass, field
from typing import Any, List, Optional

logger = logging.getLogger(__name__)


@dataclass
class ExtractResult:
    kind: str  # excel | csv | doc | pdf | txt
    full_text: str
    excel_rows: Optional[List[dict[str, Any]]] = None
    sheet_names: List[str] = field(default_factory=list)


def extract_from_bytes(file_bytes: bytes, filename: str) -> ExtractResult:
    ext = os.path.splitext(filename or "")[1].lower()

    if ext in (".txt", ".md"):
        text = sanitize_extracted_text(_decode_text(file_bytes))
        return ExtractResult(kind="txt", full_text=text, excel_rows=None)

    if ext == ".csv":
        return _extract_csv(file_bytes)

    if ext in (".xlsx", ".xls"):
        return _extract_excel(file_bytes, ext)

    if ext in (".docx", ".doc"):
        return ExtractResult(kind="doc", full_text=sanitize_extracted_text(_extract_docx(file_bytes)), excel_rows=None)

    if ext == ".pdf":
        return ExtractResult(kind="pdf", full_text=sanitize_extracted_text(_extract_pdf(file_bytes)), excel_rows=None)

    raise ValueError(f"不支持的扩展名: {ext}")


def _decode_text(b: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "latin-1"):
        try:
            return b.decode(enc)
        except Exception:
            continue
    return b.decode("utf-8", errors="replace")


def _extract_csv(file_bytes: bytes) -> ExtractResult:
    text = _decode_text(file_bytes)
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for r in reader:
        row = {
            k: (sanitize_extracted_text(v) if isinstance(v, str) else v)
            for k, v in r.items()
            if k
        }
        if any(str(v).strip() for v in row.values() if v is not None):
            rows.append(row)
    lines = []
    for row in rows:
        parts = [f"{k}: {sanitize_extracted_text(v)}" for k, v in row.items() if v is not None and str(v).strip()]
        lines.append(" | ".join(parts))
    return ExtractResult(
        kind="csv",
        full_text=sanitize_extracted_text("\n".join(lines)),
        excel_rows=rows,
        sheet_names=["csv"],
    )


def _extract_excel(file_bytes: bytes, ext: str) -> ExtractResult:
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("请安装 openpyxl") from e

    if ext == ".xls":
        try:
            import xlrd  # type: ignore
        except ImportError:
            raise RuntimeError("解析 .xls 需安装 xlrd：pip install xlrd")

        book = xlrd.open_workbook(file_contents=file_bytes)
        all_rows: List[dict] = []
        lines: List[str] = []
        sheet_names: List[str] = []
        for sheet in book.sheets():
            sheet_names.append(sheet.name)
            header = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
            for r in range(1, sheet.nrows):
                row = {}
                for c in range(sheet.ncols):
                    key = header[c] if c < len(header) else f"col{c}"
                    val = sheet.cell_value(r, c)
                    row[key] = sanitize_extracted_text(val)
                if any(str(v).strip() for v in row.values() if v is not None):
                    all_rows.append(row)
                    lines.append(
                        " | ".join(
                            f"{k}: {sanitize_extracted_text(v)}"
                            for k, v in row.items()
                            if v is not None and str(v).strip()
                        )
                    )
        return ExtractResult(
            kind="excel",
            full_text="\n".join(lines),
            excel_rows=all_rows,
            sheet_names=sheet_names,
        )

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    all_rows: List[dict] = []
    lines: List[str] = []
    sheet_names: List[str] = []
    for ws in wb.worksheets:
        sheet_names.append(ws.title)
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            continue
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header_row)]
        for row in rows_iter:
            if row is None:
                continue
            d = {}
            for i, cell in enumerate(row):
                key = header[i] if i < len(header) else f"col{i}"
                d[key] = sanitize_extracted_text(cell)
            if any(str(v).strip() for v in d.values() if v is not None):
                all_rows.append(d)
                lines.append(
                        " | ".join(
                            f"{k}: {sanitize_extracted_text(v)}"
                            for k, v in d.items()
                            if v is not None and str(v).strip()
                        )
                )
    wb.close()
    return ExtractResult(
        kind="excel",
        full_text=sanitize_extracted_text("\n".join(lines)),
        excel_rows=all_rows,
        sheet_names=sheet_names,
    )


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("请安装 python-docx") from e

    doc = Document(io.BytesIO(file_bytes))
    parts: List[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(sanitize_extracted_text(t))
    return "\n\n".join(parts)


def _extract_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        if text.strip():
            return sanitize_extracted_text(text)
    except ImportError:
        logger.warning("pdfplumber 未安装，尝试 PyPDF2")
    except Exception as e:
        logger.warning(f"pdfplumber 解析失败: {e}")

    try:
        import PyPDF2

        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
    except Exception as e:
        raise RuntimeError(f"PDF 解析失败: {e}") from e
    return sanitize_extracted_text(text)


_URL_RE = re.compile(r"https?://\S+|www\.\S+", re.IGNORECASE)
_TAG_RE = re.compile(r"(?:分类|标签|tag|category)\s*[:：]\s*[^\n]+", re.IGNORECASE)
_NOISE_CHARS_RE = re.compile(r"[`~^*_#]{2,}")
_MULTI_BLANK_RE = re.compile(r"\n{3,}")
_MULTI_SPACE_RE = re.compile(r"[ \t]{2,}")


def sanitize_extracted_text(value: Any) -> str:
    """
    统一清洗提取文本，避免无关链接/分类标签/格式噪声污染后续题库抽取与向量化。
    """
    text = str(value or "")
    if not text.strip():
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _URL_RE.sub("", text)
    text = _TAG_RE.sub("", text)
    text = _NOISE_CHARS_RE.sub(" ", text)
    text = text.replace("|", " ")
    text = _MULTI_SPACE_RE.sub(" ", text)
    text = _MULTI_BLANK_RE.sub("\n\n", text)
    lines = [ln.strip(" -\t") for ln in text.split("\n")]
    lines = [ln for ln in lines if ln and ln not in ("-", "--", "—", "——")]
    return "\n".join(lines).strip()
