import io
import logging
import os
import urllib.parse
from datetime import datetime

from fastapi import APIRouter, UploadFile, File, HTTPException, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from config.db_conf import get_db
from crud.resume import create_resume, update_resume_analysis, update_resume_optimization, get_user_resumes, delete_resume, get_resume_by_id
from models.users import User
from schemas.resume import GenerateResumeRequest
from services.resume_services import analyze_resume, optimize_resume
from services.kb.rag_plan1 import run_full_pipeline
from utils.auth import get_current_user
from utils.response import success_response

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/resume", tags=["resume"])

UPLOAD_DIR = "uploads/resumes"
os.makedirs(UPLOAD_DIR, exist_ok=True)


def _resume_record_to_docx_payload(resume) -> dict:
    """优先使用优化后的结构化数据，否则使用分析结果。"""
    if getattr(resume, "optimized_basic_info", None):
        return {
            "basic_info": resume.optimized_basic_info or {},
            "education": resume.optimized_education or [],
            "experience": resume.optimized_experience or [],
            "projects": resume.optimized_projects or [],
            "skills": resume.optimized_skills or [],
        }
    return {
        "basic_info": resume.basic_info or {},
        "education": resume.education or [],
        "experience": resume.experience or [],
        "projects": resume.projects or [],
        "skills": resume.skills or [],
    }


def _resume_payload_has_content(data: dict) -> bool:
    b = data.get("basic_info") or {}
    if isinstance(b, dict) and (b.get("name") or b.get("email") or b.get("phone")):
        return True
    for key in ("education", "experience", "projects", "skills"):
        v = data.get(key)
        if v and isinstance(v, list) and len(v) > 0:
            return True
    return False


@router.post("/upload", summary="上传简历文件（不分析）")
async def upload_resume(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
):
    allowed_types = [
        "application/pdf",
        "application/msword",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ]
    if file.content_type not in allowed_types:
        raise HTTPException(status_code=400, detail="仅支持 PDF 或 Word 格式")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    original_name = os.path.basename(file.filename or "upload.bin").replace(" ", "_")
    file_path = os.path.join(UPLOAD_DIR, f"{timestamp}_{original_name}")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="文件为空")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件超过 20MB")
    with open(file_path, "wb") as f:
        f.write(content)

    return success_response(
        message="上传成功",
        data={
            "file_id": int(datetime.now().timestamp()),
            "file_name": original_name,
            "saved_path": file_path,
            "user_id": user.id,
        }
    )


@router.get("/download/{resume_id}", summary="按简历 ID 下载 Word 或 PDF（优先优化版）")
async def download_resume_by_id(
    resume_id: int,
    format: str = Query("docx", description="docx 或 pdf"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    from utils.pdf_export import resume_data_to_pdf

    fmt = (format or "docx").lower()
    if fmt not in ("docx", "pdf"):
        raise HTTPException(status_code=400, detail="format 仅支持 docx 或 pdf")

    resume = await get_resume_by_id(db, resume_id)
    if not resume or resume.user_id != user.id:
        raise HTTPException(status_code=404, detail="简历不存在或无权访问")

    payload = _resume_record_to_docx_payload(resume)
    if not _resume_payload_has_content(payload):
        raise HTTPException(status_code=400, detail="简历数据不完整，请先完成分析、优化或生成")

    name = (payload.get("basic_info") or {}).get("name", "简历")
    if fmt == "pdf":
        buffer = resume_data_to_pdf(payload)
        encoded_filename = urllib.parse.quote(f"简历_{name}.pdf")
        media = "application/pdf"
    else:
        buffer = generate_resume_docx(payload)
        encoded_filename = urllib.parse.quote(f"简历_{name}.docx")
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    return StreamingResponse(
        buffer,
        media_type=media,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post("/analyze", summary="上传并分析简历（智能扫描）")
async def analyze_resume_endpoint(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    allowed_types = [
        "application/pdf", "application/msword",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "text/plain",
    ]
    safe_name = os.path.basename(file.filename or "upload.bin").replace(" ", "_").replace("(", "_").replace(")", "_")
    if file.content_type not in allowed_types and not safe_name.lower().endswith((".pdf", ".doc", ".docx", ".txt")):
        raise HTTPException(status_code=400, detail="仅支持 PDF、Word 或 TXT 格式")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saved_path = os.path.join(UPLOAD_DIR, f"{timestamp}_{safe_name}")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="文件为空")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件超过 20MB")
    with open(saved_path, "wb") as f:
        f.write(content)

    try:
        resume_text = await extract_text_from_file(saved_path, safe_name)
        if not resume_text or not resume_text.strip():
            raise HTTPException(status_code=400, detail="无法从文件中提取文本内容")

        resume_data = await analyze_resume(resume_text)

        resume_record = await create_resume(db, user.id, safe_name, saved_path)
        await update_resume_analysis(db, resume_record.id, resume_data)
        await run_full_pipeline(
            file_path=saved_path,
            file_info={
                "filename": safe_name,
                "file_type": "resume",
                "status": "active",
                "job_id": None,
            },
            db=db,
        )

        logger.info(f"[简历分析] 用户 {user.username} 完成，ID={resume_record.id}")

        return success_response(
            message="分析成功",
            data={"resume_id": resume_record.id, **resume_data}
        )
    except HTTPException:
        try:
            if os.path.isfile(saved_path):
                os.remove(saved_path)
        except OSError:
            pass
        raise
    except Exception:
        try:
            if os.path.isfile(saved_path):
                os.remove(saved_path)
        except OSError:
            pass
        raise


async def extract_text_from_file(file_path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()

    elif ext in ['.doc', '.docx']:
        try:
            from docx import Document
            doc = Document(file_path)
            return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            raise HTTPException(status_code=500, detail="服务器缺少文档处理依赖")

    elif ext == '.pdf':
        try:
            import PyPDF2
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                return ''.join(page.extract_text() + '\n' for page in reader.pages)
        except ImportError:
            pass
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                return ''.join((p.extract_text() or '') + '\n' for p in pdf.pages)
        except ImportError:
            pass
        raise HTTPException(status_code=500, detail="服务器缺少 PDF 处理依赖，请安装 PyPDF2 或 pdfplumber")

    raise HTTPException(status_code=400, detail=f"不支持的文件类型：{ext}")


@router.post("/optimize-data", summary="AI 优化简历（传递数据版本）")
async def optimize_resume_with_data(
    resume_data: dict,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if not resume_data or not isinstance(resume_data, dict):
        raise HTTPException(status_code=400, detail="无效的简历数据格式")

    optimized_data = await optimize_resume(resume_data)

    resume_id = resume_data.get('resume_id')
    if resume_id:
        await update_resume_optimization(db, resume_id, optimized_data)

    return success_response(
        message="优化成功",
        data={
            "optimized_resume": optimized_data,
            "optimized_text": generate_resume_text(optimized_data),
        }
    )


def generate_resume_text(resume_data: dict) -> str:
    text = "AI 优化版简历\n" + "=" * 50 + "\n\n"
    basic = resume_data.get('basic_info', {})
    text += f"【基本信息】\n姓名：{basic.get('name', '')}\n邮箱：{basic.get('email', '')}\n电话：{basic.get('phone', '')}\n\n"

    text += "【教育背景】\n"
    for edu in resume_data.get('education', []):
        text += f"{edu.get('school', '')} | {edu.get('degree', '')} | {edu.get('major', '')}\n"
        text += f"时间：{edu.get('start_date', '')} - {edu.get('end_date', '')}\n"
        if edu.get('honors'):
            text += f"荣誉：{edu.get('honors')}\n"
        text += "\n"

    text += "【工作经历】\n"
    for exp in resume_data.get('experience', []):
        text += f"{exp.get('company', '')} | {exp.get('position', '')}\n"
        text += f"时间：{exp.get('start_date', '')} - {exp.get('end_date', '')}\n"
        ach = exp.get("achievements", "")
        if isinstance(ach, list):
            ach_str = "；".join(str(x) for x in ach)
        else:
            ach_str = str(ach or "")
        text += f"主要成就：\n{ach_str}\n\n"

    text += "【项目经验】\n"
    for proj in resume_data.get('projects', []):
        text += f"项目名称：{proj.get('name', '')}\n担任角色：{proj.get('role', '')}\n项目描述：\n{proj.get('description', '')}\n\n"

    text += "【技能清单】\n" + "、".join(resume_data.get('skills', [])) + "\n"
    return text


@router.get("/history", summary="获取简历历史")
async def get_resume_history(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    resumes = await get_user_resumes(db, user.id)
    resume_list = [
        {
            "id": r.id,
            "file_name": r.file_name,
            "status": r.status,
            "basic_info": r.basic_info or {},
            "created_at": str(r.credited_at) if hasattr(r, 'credited_at') else "未知时间",
        }
        for r in resumes
    ]
    return success_response(message="获取成功", data={"resumes": resume_list})


@router.delete("/{resume_id}", summary="删除简历")
async def delete_resume_endpoint(
    resume_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    resume = await get_resume_by_id(db, resume_id)
    if not resume:
        raise HTTPException(status_code=404, detail="简历不存在")
    if resume.user_id != user.id:
        raise HTTPException(status_code=403, detail="无权删除此简历")

    if os.path.exists(resume.file_path):
        try:
            os.remove(resume.file_path)
        except Exception:
            pass

    success = await delete_resume(db, resume_id)
    if not success:
        raise HTTPException(status_code=500, detail="删除失败")
    return success_response(message="删除成功")


@router.post("/generate", summary="AI 生成简历")
async def generate_resume_endpoint(
    user_info: GenerateResumeRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    for field in ('name', 'email'):
        if not getattr(user_info, field, None):
            raise HTTPException(status_code=400, detail=f"缺少必要字段：{field}")

    from config.llm_conf import DEEPSEEK_MODEL, call_llm
    from services.LLM_services import _parse_llm_json_dict

    prompt = f"""请根据以下基本信息，生成一份完整的专业简历。

基本信息：
- 姓名：{user_info.name}
- 性别：{user_info.gender or '未提供'}
- 邮箱：{user_info.email}
- 电话：{user_info.phone}
- 求职意向：{user_info.job_target or '未提供'}
- 工作年限：{user_info.experience_years or '未提供'}

教育背景：
{format_list_for_prompt(user_info.education)}

工作经历：
{format_list_for_prompt(user_info.experience)}

项目经验：
{format_list_for_prompt(user_info.projects)}

技能清单：
{', '.join(user_info.skills)}

请优化和完善这份简历，使其更加专业和有吸引力。直接返回 JSON，不要其他说明：
{{
    "basic_info": {{"name": "", "email": "", "phone": ""}},
    "education": [],
    "experience": [],
    "projects": [],
    "skills": []
}}"""

    response = await call_llm(
        prompt=prompt,
        system_prompt="你是专业的简历生成专家，擅长根据基本信息生成高质量的专业简历。",
        model=DEEPSEEK_MODEL,
    )

    generated_resume, _parse_err = _parse_llm_json_dict(response)
    if generated_resume is None:
        generated_resume = {
            "basic_info": {"name": user_info.name, "email": user_info.email, "phone": user_info.phone or ""},
            "education": list(user_info.education),
            "experience": list(user_info.experience),
            "projects": list(user_info.projects),
            "skills": list(user_info.skills),
        }

    resume_record = await create_resume(db, user.id, "generated_resume.json", "")
    await update_resume_analysis(db, resume_record.id, generated_resume)
    generated_text = generate_resume_text(generated_resume)
    generated_path = os.path.join(UPLOAD_DIR, f"generated_{resume_record.id}.txt")
    with open(generated_path, "w", encoding="utf-8") as f:
        f.write(generated_text)
    await run_full_pipeline(
        file_path=generated_path,
        file_info={
            "filename": f"generated_resume_{resume_record.id}.txt",
            "file_type": "resume",
            "status": "active",
            "job_id": None,
        },
        db=db,
    )

    return success_response(
        message="生成成功",
        data={"resume_id": resume_record.id, **generated_resume}
    )


def format_list_for_prompt(items: list) -> str:
    if not items:
        return "无"
    result = ""
    for i, item in enumerate(items, 1):
        if isinstance(item, dict):
            result += f"{i}. " + ", ".join(f"{k}: {v}" for k, v in item.items()) + "\n"
        else:
            result += f"{i}. {item}\n"
    return result


def generate_resume_docx(resume_data: dict) -> io.BytesIO:
    """
    生成 Word 简历：统一中文字体（eastAsia）、段落行距与段后距，避免西文默认字体导致中文显示发虚、排版拥挤。
    """
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    # Windows 常见中文字体；若系统无「微软雅黑」可改为「宋体」
    FONT_CN = "微软雅黑"
    FONT_CN_TITLE = "微软雅黑"

    def set_run_east_asia_font(run, font_name: str) -> None:
        run.font.name = font_name
        r = run._element
        r_pr = r.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:ascii"), font_name)
        r_fonts.set(qn("w:hAnsi"), font_name)
        r_fonts.set(qn("w:eastAsia"), font_name)

    def fmt_para(paragraph, space_after_pt=6, line_multiple=1.2, first_line_indent_pt=0):
        pf = paragraph.paragraph_format
        pf.space_after = Pt(space_after_pt)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_multiple
        if first_line_indent_pt:
            pf.first_line_indent = Pt(first_line_indent_pt)

    def add_run(paragraph, text: str, *, size_pt=10.5, bold=False, color_rgb=None, font=FONT_CN):
        if text is None:
            text = ""
        r = paragraph.add_run(str(text))
        set_run_east_asia_font(r, font)
        r.font.size = Pt(size_pt)
        r.bold = bold
        if color_rgb:
            r.font.color.rgb = RGBColor(*color_rgb)
        return r

    def add_section_title(doc, title: str):
        p = doc.add_paragraph()
        fmt_para(p, space_after_pt=8, line_multiple=1.0)
        add_run(p, title, size_pt=12, bold=True, color_rgb=(0x1E, 0x40, 0xAF), font=FONT_CN_TITLE)
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "2563EB")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.5)

    basic = resume_data.get("basic_info", {}) or {}

    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_para(name_para, space_after_pt=4, line_multiple=1.0)
    add_run(name_para, basic.get("name", "简历") or "简历", size_pt=26, bold=True, color_rgb=(0x0F, 0x17, 0x2A), font=FONT_CN_TITLE)

    contact_bits = []
    if basic.get("email"):
        contact_bits.append(f"邮箱：{basic['email']}")
    if basic.get("phone"):
        contact_bits.append(f"电话：{basic['phone']}")
    if basic.get("job_target"):
        contact_bits.append(f"求职意向：{basic['job_target']}")
    if contact_bits:
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_para(contact_para, space_after_pt=14, line_multiple=1.15)
        add_run(contact_para, "  ·  ".join(contact_bits), size_pt=10, color_rgb=(0x64, 0x74, 0x8B))

    education = resume_data.get("education", []) or []
    if education:
        add_section_title(doc, "教育背景")
        for edu in education:
            if not isinstance(edu, dict):
                continue
            p = doc.add_paragraph()
            fmt_para(p, space_after_pt=2, line_multiple=1.2)
            add_run(p, edu.get("school", "") or "", size_pt=11, bold=True, color_rgb=(0x15, 0x23, 0x4A))
            add_run(p, f"    {edu.get('degree', '')} · {edu.get('major', '')}", size_pt=10.5, color_rgb=(0x47, 0x55, 0x69))
            t = doc.add_paragraph()
            fmt_para(t, space_after_pt=8, line_multiple=1.2)
            add_run(t, f"{edu.get('start_date', '')} — {edu.get('end_date', '至今')}", size_pt=9.5, color_rgb=(0x94, 0xA3, 0xB8))
            if edu.get("honors"):
                h = doc.add_paragraph()
                fmt_para(h, space_after_pt=6, line_multiple=1.25, first_line_indent_pt=0)
                add_run(h, f"荣誉：{edu['honors']}", size_pt=10, color_rgb=(0x51, 0x65, 0x85))

    experience = resume_data.get("experience", []) or []
    if experience:
        add_section_title(doc, "工作经历")
        for exp in experience:
            if not isinstance(exp, dict):
                continue
            p = doc.add_paragraph()
            fmt_para(p, space_after_pt=2, line_multiple=1.2)
            add_run(p, exp.get("company", "") or "", size_pt=11, bold=True, color_rgb=(0x15, 0x23, 0x4A))
            add_run(p, f"    {exp.get('position', '')}", size_pt=10.5, color_rgb=(0x47, 0x55, 0x69))
            t = doc.add_paragraph()
            fmt_para(t, space_after_pt=6, line_multiple=1.2)
            add_run(t, f"{exp.get('start_date', '')} — {exp.get('end_date', '至今')}", size_pt=9.5, color_rgb=(0x94, 0xA3, 0xB8))
            achievements = exp.get("achievements") or []
            if isinstance(achievements, list):
                for ach in achievements:
                    if not ach:
                        continue
                    bp = doc.add_paragraph()
                    fmt_para(bp, space_after_pt=4, line_multiple=1.25, first_line_indent_pt=0)
                    add_run(bp, f"•  {str(ach).strip()}", size_pt=10.5, color_rgb=(0x33, 0x41, 0x55))
            elif achievements:
                bp = doc.add_paragraph()
                fmt_para(bp, space_after_pt=6, line_multiple=1.25)
                add_run(bp, f"•  {str(achievements).strip()}", size_pt=10.5, color_rgb=(0x33, 0x41, 0x55))

    projects = resume_data.get("projects", []) or []
    if projects:
        add_section_title(doc, "项目经验")
        for proj in projects:
            if not isinstance(proj, dict):
                continue
            p = doc.add_paragraph()
            fmt_para(p, space_after_pt=2, line_multiple=1.2)
            add_run(p, proj.get("name", "") or "", size_pt=11, bold=True, color_rgb=(0x15, 0x23, 0x4A))
            if proj.get("role"):
                add_run(p, f"（{proj['role']}）", size_pt=10, color_rgb=(0x71, 0x85, 0x9A))
            if proj.get("description"):
                d = doc.add_paragraph()
                fmt_para(d, space_after_pt=10, line_multiple=1.3, first_line_indent_pt=0)
                add_run(d, str(proj["description"]).strip(), size_pt=10.5, color_rgb=(0x33, 0x41, 0x55))

    skills = resume_data.get("skills", []) or []
    if skills:
        add_section_title(doc, "技能清单")
        skills_p = doc.add_paragraph()
        fmt_para(skills_p, space_after_pt=6, line_multiple=1.35)
        text = "、".join(skills) if isinstance(skills, list) else str(skills)
        add_run(skills_p, text, size_pt=10.5, color_rgb=(0x33, 0x41, 0x55))

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


@router.post("/download/docx", summary="将简历数据生成并下载为 Word 或 PDF 文档")
async def download_resume_docx(
    resume_data: dict,
    format: str = Query("docx", description="docx 或 pdf"),
    user: User = Depends(get_current_user),
):
    from utils.pdf_export import resume_data_to_pdf

    fmt = (format or "docx").lower()
    if fmt not in ("docx", "pdf"):
        raise HTTPException(status_code=400, detail="format 仅支持 docx 或 pdf")

    if not resume_data or not isinstance(resume_data, dict):
        raise HTTPException(status_code=400, detail="无效的简历数据")

    name = resume_data.get('basic_info', {}).get('name', '简历')
    if fmt == "pdf":
        buffer = resume_data_to_pdf(resume_data)
        encoded_filename = urllib.parse.quote(f"简历_{name}.pdf")
        media = "application/pdf"
    else:
        buffer = generate_resume_docx(resume_data)
        encoded_filename = urllib.parse.quote(f"简历_{name}.docx")
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    return StreamingResponse(
        buffer,
        media_type=media,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )


@router.post("/save-generated", summary="保存生成的简历")
async def save_generated_resume(
    resume_data: dict,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    resume_record = await create_resume(db, user.id, "generated_resume.json", "")
    await update_resume_analysis(db, resume_record.id, resume_data)
    generated_text = generate_resume_text(resume_data)
    generated_path = os.path.join(UPLOAD_DIR, f"saved_generated_{resume_record.id}.txt")
    with open(generated_path, "w", encoding="utf-8") as f:
        f.write(generated_text)
    await run_full_pipeline(
        file_path=generated_path,
        file_info={
            "filename": f"saved_generated_resume_{resume_record.id}.txt",
            "file_type": "resume",
            "status": "active",
            "job_id": None,
        },
        db=db,
    )
    return success_response(message="保存成功", data={"resume_id": resume_record.id})
